"""Bangkitkan dataset sintetis dalam FORMAT WORKBOOK BGG yang sebenarnya.

Dipakai untuk menguji seluruh rantai end-to-end - termasuk resolver header
merged bertingkat - tanpa data produksi. Header ditulis dengan struktur merge
yang sama seperti workbook asli, sehingga resolver benar-benar teruji dan bukan
dilewati.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import ezdxf
import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

RNG = np.random.default_rng(5015)
X0, Y0 = 359_000.0, 9_594_000.0

SEAMS = [("S11", 22.0, 1.60), ("S10A", 45.0, 9.90), ("S10B", 58.0, 11.80)]
QUALITY = {
    "S11":  dict(tm=41.0, im=17.0, ash=13.5, ts=0.25, cv_adb=4700, rd=1.38),
    "S10A": dict(tm=42.3, im=17.6, ash=10.8, ts=0.20, cv_adb=4815, rd=1.36),
    "S10B": dict(tm=43.7, im=18.6, ash=8.1,  ts=0.31, cv_adb=5022, rd=1.35),
}
DIP_DEG, DIP_AZ = 4.0, 115.0

LITHOLOGY_LIBRARY = [
    ("CO", "Coal Undifferentiated"), ("C3", "Coal, 40-60% bright"),
    ("C5", "Coal, <10% bright"), ("CS", "Clay Stone"), ("SS", "Sandstone"),
    ("XC", "Carbonaceous CS"), ("SO", "Soil"), ("KL", "Core Loss"),
    ("ZC", "Coaly Clystone"), ("LG", "Lignite"),
]


def topo_rl(x, y):
    dx, dy = (x - X0) / 1000.0, (y - Y0) / 1000.0
    return 66.0 + 6.0 * np.sin(1.7 * dx) * np.cos(1.3 * dy) + 3.5 * np.sin(3.1 * dy) - 2.0 * dx


def seam_geometry(name, x, y):
    depth0, mean_t = next((d, t) for n, d, t in SEAMS if n == name)
    az = np.radians(DIP_AZ)
    down_dip = (x - X0) * np.sin(az) + (y - Y0) * np.cos(az)
    dx, dy = (x - X0) / 1000.0, (y - Y0) / 1000.0
    roof_depth = depth0 + down_dip * np.tan(np.radians(DIP_DEG)) - 2.0 * np.sin(2.2 * dx)
    thickness = max(0.3, mean_t + 0.8 * np.sin(1.9 * dx) * np.cos(2.4 * dy))
    return float(roof_depth), float(thickness)


def _write_block(ws, row, col, value, rowspan=1, colspan=1):
    ws.cell(row=row, column=col, value=value)
    if rowspan > 1 or colspan > 1:
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row + rowspan - 1, end_column=col + colspan - 1)


def build_collar_sheet(wb, hole_id, east, north, rl, td, log_depth):
    ws = wb.create_sheet("Collar")
    _write_block(ws, 2, 1, "Hole ID", rowspan=3)
    _write_block(ws, 2, 2, "Collar", colspan=7)
    _write_block(ws, 3, 2, "Coordinate", colspan=3)
    for offset, name in enumerate(("East", "North", "RL")):
        _write_block(ws, 4, 2 + offset, name)
    _write_block(ws, 3, 5, "Survey", colspan=4)
    for offset, name in enumerate(("Date", "Company", "Method", "Surveyor")):
        _write_block(ws, 4, 5 + offset, name)
    _write_block(ws, 2, 9, "Drilling", colspan=8)
    _write_block(ws, 3, 16, "Max Depth", rowspan=2)
    _write_block(ws, 2, 17, "Geophysical", colspan=7)
    _write_block(ws, 3, 17, "Company", rowspan=2)
    _write_block(ws, 3, 18, "Logger", rowspan=2)
    _write_block(ws, 3, 19, "Geophysical Log", colspan=4)
    for offset, name in enumerate(("GM", "DN", "CAL", "RS")):
        _write_block(ws, 4, 19 + offset, name)
    _write_block(ws, 3, 23, "Max Depth", rowspan=2)

    values = {1: hole_id, 2: round(east, 3), 3: round(north, 3), 4: round(rl, 3),
              6: "BGG", 7: "TS", 16: td, 17: "REC", 18: "DIDI",
              19: "v", 20: "v", 21: "v", 22: "x", 23: round(log_depth, 2)}
    for col, value in values.items():
        ws.cell(row=5, column=col, value=value)
    return ws


def build_sll_sheet(wb, name, hole_id, intervals, log_basis):
    ws = wb.create_sheet(name)
    ws.cell(row=2, column=1, value="PT. BUDI GEMA GEMPITA")
    ws.cell(row=3, column=1, value="Sedimentary Lithology Log")
    ws.cell(row=5, column=1, value="Prospect Area")
    ws.cell(row=5, column=3, value=": BGG Block Lawai 1")

    _write_block(ws, 7, 1, "Hole_ID", rowspan=3)
    _write_block(ws, 7, 2, "Drill Detail", colspan=15)
    for offset, label in enumerate(("Shift", "Log Basis", "Bit Type", "Hole Dia", "Hole Type")):
        _write_block(ws, 8, 2 + offset, label, rowspan=2)
    _write_block(ws, 8, 7, "Depth", colspan=2)
    _write_block(ws, 9, 7, "From")
    _write_block(ws, 9, 8, "To")
    _write_block(ws, 8, 9, "Core Length", rowspan=2)
    _write_block(ws, 8, 10, "Core Recov", rowspan=2)
    _write_block(ws, 7, 17, "Lithology", colspan=6)
    _write_block(ws, 8, 17, "Lith", rowspan=2)
    _write_block(ws, 8, 18, "Lith Qualifier", rowspan=2)
    _write_block(ws, 8, 19, "Seam", rowspan=2)
    _write_block(ws, 8, 20, "Ply", rowspan=2)
    for col in range(1, 21):
        ws.cell(row=10, column=col, value="-")

    for index, item in enumerate(intervals):
        row = 11 + index
        ws.cell(row=row, column=1, value=hole_id if index == 0 else None)
        ws.cell(row=row, column=3, value=log_basis)
        ws.cell(row=row, column=6, value=item["hole_type"])
        ws.cell(row=row, column=7, value=round(item["from"], 3))
        ws.cell(row=row, column=8, value=round(item["to"], 3))
        if item.get("core_recovery") is not None:
            ws.cell(row=row, column=10, value=item["core_recovery"])
        ws.cell(row=row, column=17, value=item["lith"])
        ws.cell(row=row, column=19, value=item.get("seam"))
    return ws


def build_bhc_sheet(wb, hole_id, east, north, rl, td, log_depth, seam_rows):
    ws = wb.create_sheet("BHC")
    ws.cell(row=2, column=1, value="PT. BUDI GEMA GEMPITA")
    ws.cell(row=3, column=1, value="BORE HOLE COMPLETION")
    _write_block(ws, 7, 1, "NO", rowspan=3)
    _write_block(ws, 7, 2, "HOLE NUMBER", colspan=2)
    _write_block(ws, 8, 2, "PLAN", rowspan=2)
    _write_block(ws, 8, 3, "ACTUAL", rowspan=2)
    _write_block(ws, 7, 9, "COORDINATE BY GPS", colspan=3)
    for offset, label in enumerate(("EASTING", "NORTHING", "ELV.")):
        _write_block(ws, 8, 9 + offset, label, rowspan=2)
    _write_block(ws, 7, 12, "COORDINATE BY TS", colspan=3)
    for offset, label in enumerate(("EASTING", "NORTHING", "ELV.")):
        _write_block(ws, 8, 12 + offset, label, rowspan=2)
    _write_block(ws, 7, 22, "CORE LENGTH ( M )", rowspan=3)
    _write_block(ws, 7, 23, "CORE REC. ( % )", rowspan=3)
    _write_block(ws, 7, 24, "TOTAL DEPTH (m)", colspan=2)
    _write_block(ws, 8, 24, "DRILLING", rowspan=2)
    _write_block(ws, 8, 25, "LOGGING", rowspan=2)
    _write_block(ws, 7, 32, "PICKING FROM", rowspan=3)
    _write_block(ws, 7, 33, "RECONCILE", colspan=3)
    for offset, label in enumerate(("FROM", "TO", "THICK")):
        _write_block(ws, 8, 33 + offset, label, rowspan=2)
    _write_block(ws, 7, 36, "COAL REC ( % )", rowspan=3)
    _write_block(ws, 7, 37, "SEAM", rowspan=3)

    ws.cell(row=10, column=1, value=1)
    ws.cell(row=10, column=2, value=hole_id)
    ws.cell(row=10, column=3, value=hole_id)
    # GPS handheld sengaja meleset dari total station, seperti data nyata.
    ws.cell(row=10, column=9, value=round(east + RNG.normal(0, 3), 3))
    ws.cell(row=10, column=10, value=round(north + RNG.normal(0, 3), 3))
    ws.cell(row=10, column=11, value=round(rl + RNG.normal(0, 4), 3))
    ws.cell(row=10, column=12, value=round(east, 3))
    ws.cell(row=10, column=13, value=round(north, 3))
    ws.cell(row=10, column=14, value=round(rl, 3))
    ws.cell(row=10, column=24, value=td)
    ws.cell(row=10, column=25, value=round(log_depth, 2))
    for index, item in enumerate(seam_rows):
        row = 10 + index
        ws.cell(row=row, column=32, value="LSD")
        ws.cell(row=row, column=33, value=round(item["from"], 3))
        ws.cell(row=row, column=34, value=round(item["to"], 3))
        ws.cell(row=row, column=35, value=round(item["to"] - item["from"], 3))
        ws.cell(row=row, column=37, value=item["seam"])
    return ws


def build_sampling_sheet(wb, hole_id, samples):
    ws = wb.create_sheet("Sampling")
    _write_block(ws, 2, 2, "No", rowspan=2)
    _write_block(ws, 2, 3, "Hole Id", rowspan=2)
    _write_block(ws, 2, 4, "Sample Position", rowspan=2)
    _write_block(ws, 2, 5, "Sample Number", rowspan=2)
    _write_block(ws, 2, 6, "Drilling Depth", colspan=3)
    for offset, label in enumerate(("From", "To", "Length")):
        _write_block(ws, 3, 6 + offset, label)
    _write_block(ws, 2, 9, "Adjusted Depth", colspan=3)
    for offset, label in enumerate(("From", "To", "Length")):
        _write_block(ws, 3, 9 + offset, label)
    _write_block(ws, 2, 12, "Sampling Date", rowspan=2)
    _write_block(ws, 2, 13, "Seam", rowspan=2)
    _write_block(ws, 2, 14, "Remarks", rowspan=2)

    for index, sample in enumerate(samples):
        row = 5 + index
        ws.cell(row=row, column=2, value=index + 1)
        ws.cell(row=row, column=3, value=hole_id if index == 0 else None)
        ws.cell(row=row, column=4, value=sample["position"])
        ws.cell(row=row, column=5, value=sample["id"])
        ws.cell(row=row, column=6, value=round(sample["from"], 3))
        ws.cell(row=row, column=7, value=round(sample["to"], 3))
        ws.cell(row=row, column=8, value=round(sample["to"] - sample["from"], 3))
        ws.cell(row=row, column=9, value=round(sample["from"], 3))
        ws.cell(row=row, column=10, value=round(sample["to"], 3))
        ws.cell(row=row, column=11, value=round(sample["to"] - sample["from"], 3))
        ws.cell(row=row, column=13, value=sample["seam"])
    return ws


def build_library_sheet(wb):
    ws = wb.create_sheet("Library SLL")
    ws.cell(row=1, column=1, value="Shift")
    ws.cell(row=1, column=4, value="Lithology")
    ws.cell(row=1, column=7, value="Qualifers")
    ws.cell(row=1, column=10, value="Grain")
    ws.cell(row=1, column=13, value="Sediment Structure")
    ws.cell(row=1, column=16, value="Weathering")
    for index, (code, description) in enumerate(LITHOLOGY_LIBRARY):
        ws.cell(row=2 + index, column=4, value=code)
        ws.cell(row=2 + index, column=5, value=description)
    ws.cell(row=2, column=1, value="D")
    ws.cell(row=2, column=2, value="Day")
    return ws


def build_hole(path: Path, hole_id, east, north, cored: bool):
    rl = float(topo_rl(np.array(east), np.array(north)))
    intervals, seam_rows, samples = [], [], []
    cursor = 0.0
    sample_index = 1

    for name, _, _ in SEAMS:
        roof, thickness = seam_geometry(name, east, north)
        if roof <= cursor:
            continue
        intervals.append({"from": cursor, "to": roof, "lith": "CS",
                          "seam": None, "hole_type": "FC" if cored else "OH"})
        floor = roof + thickness
        parting_at = roof + thickness * 0.45
        segments = [
            (roof, parting_at, "C3", name),
            (parting_at, parting_at + 0.08, "XC", None),
            (parting_at + 0.08, floor, "C3", name),
        ]
        if cored and name == "S10B":
            mid = parting_at + 0.08 + (floor - parting_at - 0.08) * 0.4
            segments = [
                (roof, parting_at, "C3", name),
                (parting_at, parting_at + 0.08, "XC", None),
                (parting_at + 0.08, mid, "C3", name),
                (mid, mid + 0.06, "KL", name),          # core loss di dalam seam
                (mid + 0.06, floor, "C3", name),
            ]
        for seg_from, seg_to, lith, seam in segments:
            intervals.append({
                "from": seg_from, "to": seg_to, "lith": lith, "seam": seam,
                "hole_type": "FC" if cored else "OH",
                "core_recovery": round(float(RNG.uniform(95, 100)), 1) if cored else None,
            })
        seam_rows.append({"from": roof, "to": floor, "seam": name})
        cursor = floor

        if cored:
            for position, (s_from, s_to) in (
                ("TOP-COAL", (roof, roof + 0.10)),
                ("BODY-COAL", (roof + 0.10, floor - 0.10)),
                ("BOTTOM-COAL", (floor - 0.10, floor)),
            ):
                samples.append({
                    "id": f"{hole_id}-SPL_{sample_index:02d}", "position": position,
                    "from": s_from, "to": s_to, "seam": name,
                })
                sample_index += 1

    td = float(np.ceil((cursor + 8.0) / 5.0) * 5.0)
    intervals.append({"from": cursor, "to": td, "lith": "SS", "seam": None,
                      "hole_type": "OH"})
    log_depth = td - 1.5

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_sll_sheet(wb, "SLL_Wellsite", hole_id,
                    [{**i, "from": round(i["from"], 1), "to": round(i["to"], 1)}
                     for i in intervals], "A")
    build_sll_sheet(wb, "SLL_Reconciled", hole_id, intervals, "X")
    build_collar_sheet(wb, hole_id, east, north, rl, td, log_depth)
    build_bhc_sheet(wb, hole_id, east, north, rl, td, log_depth, seam_rows)
    build_sampling_sheet(wb, hole_id, samples)
    build_library_sheet(wb)
    wb.save(path)
    return samples, td, log_depth


def build_las(path: Path, hole_id, td, log_depth):
    depths = np.arange(0.22, log_depth, 0.02)
    gr = 60 + 25 * np.sin(depths / 3.0) + RNG.normal(0, 6, depths.size)
    ld = 4200 + 1500 * np.cos(depths / 2.5) + RNG.normal(0, 120, depths.size)
    lines = [
        "~Version Information", "VERS.             2.0: LAS Version 2.0",
        "WRAP.             NO: One Line per depth step", "~Well Information Block",
        f"STRT.M                {depths[0]:.2f}: Start Depth",
        f"STOP.M                {depths[-1]:.2f}: Stop Depth",
        "STEP.M                0.02: Step", "NULL.                    -999:NULL VALUE",
        "COMP.                PT.BGG: COMPANY", f"WELL.                {hole_id}:",
        "FLD.                 MUARA LAWAI: FIELD", "PROV.                SOUTH SUMATERA: PROVINCE",
        "SRVC.                RescaLog: SERVICE COMPANY", "DATE.                06-04-2021: LOG DATE",
        "~Curve Information Block", "DEPT.M        :DEPTH", "GR.CPS                  :Gamma Ray",
        "CL.CPS                  :Caliper", "LD.CPS                  :Long Density",
        "SD.CPS                  :Short Density", "#DATA", "~A",
    ]
    for d, g, l in zip(depths, gr, ld):
        lines.append(f"{d:.2f} {g:.5f} {3.6:.5f} {l:.5f} {l * 4.8:.5f}")
    path.write_text("\n".join(lines))


def build_quality_csv(path: Path, records):
    pd.DataFrame(records).to_csv(path, index=False)


def build_topo_dxf(path: Path, east, north):
    doc = ezdxf.new("R2010")
    msp = doc.modelspace()
    pad = 500.0
    xs = np.linspace(east.min() - pad, east.max() + pad, 70)
    ys = np.linspace(north.min() - pad, north.max() + pad, 70)
    gx, gy = np.meshgrid(xs, ys)
    gz = topo_rl(gx, gy)
    for x, y, z in zip(gx.ravel(), gy.ravel(), gz.ravel()):
        msp.add_point((x, y, z), dxfattribs={"layer": "TOPO_SPOT"})
    doc.saveas(path)


def build(n_holes: int, cored_fraction: float, out_dir: Path) -> Path:
    workbook_dir = out_dir / "workbooks"
    las_dir = out_dir / "las"
    for directory in (workbook_dir, las_dir):
        directory.mkdir(parents=True, exist_ok=True)

    side = int(np.ceil(np.sqrt(n_holes)))
    gx, gy = np.meshgrid(np.arange(side), np.arange(side))
    # Spasi 700 m dipilih agar sel Voronoi melampaui radius measured (250 m),
    # sehingga pita Tertunjuk dan Tereka ikut terbentuk dan teruji.
    east = X0 + gx.ravel()[:n_holes] * 700.0 + RNG.normal(0, 40, n_holes)
    north = Y0 + gy.ravel()[:n_holes] * 700.0 + RNG.normal(0, 40, n_holes)
    cored = RNG.random(n_holes) < cored_fraction

    quality_records = []
    for index in range(n_holes):
        hole_id = f"DH{index // 10 + 1:02d}_{index % 10 + 1:02d}C1" if cored[index] \
            else f"DH{index // 10 + 1:02d}_{index % 10 + 1:02d}"
        samples, td, log_depth = build_hole(
            workbook_dir / f"{hole_id}.xlsx", hole_id, east[index], north[index], bool(cored[index])
        )
        build_las(las_dir / f"{hole_id}.LAS", hole_id, td, log_depth)

        if not cored[index]:
            continue
        by_seam: dict[str, list[str]] = {}
        for sample in samples:
            by_seam.setdefault(sample["seam"], []).append(sample["id"])
        for seam, ids in by_seam.items():
            base = QUALITY[seam]
            ash = max(2.0, base["ash"] + RNG.normal(0, 1.0))
            im = base["im"] + RNG.normal(0, 0.5)
            tm = base["tm"] + RNG.normal(0, 0.8)
            vm = 40.5 + RNG.normal(0, 0.6)
            cv_adb = base["cv_adb"] - 55.0 * (ash - base["ash"]) + RNG.normal(0, 50)
            quality_records.append({
                "sample_id": f"{hole_id}-{seam}-COMP", "lab_sample_id": f"PH21.{RNG.integers(1e4,9e4)}",
                "hole_id": hole_id, "seam": seam, "composite_of": ";".join(ids),
                "TM_ar": round(tm, 2), "M_adb": round(im, 2), "ASH_adb": round(ash, 2),
                "VM_adb": round(vm, 2), "FC_adb": round(100 - im - ash - vm, 2),
                "TS_adb": round(max(0.05, base["ts"] + RNG.normal(0, 0.04)), 2),
                "CV_adb": round(cv_adb), "CV_ar": round(cv_adb * (100 - tm) / (100 - im)),
                "CV_daf": round(cv_adb * 100 / (100 - im - ash)),
                "RD": round(base["rd"] + 0.0075 * (ash - base["ash"]), 3),
                "RD_basis": "air_dried", "report_no": "01221.00923",
            })

    build_quality_csv(out_dir / "quality.csv", quality_records)
    build_topo_dxf(out_dir / "topo.dxf", east, north)
    print(f"Ditulis ke {out_dir}: {n_holes} workbook ({int(cored.sum())} cored), "
          f"{n_holes} LAS, {len(quality_records)} hasil kualitas, topo.dxf")
    return out_dir


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--holes", type=int, default=25)
    ap.add_argument("--cored-fraction", type=float, default=0.45)
    ap.add_argument("--out", default="sample_data")
    args = ap.parse_args()
    build(args.holes, args.cored_fraction, Path(args.out))
