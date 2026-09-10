"""Pembaca workbook Excel dengan header gabungan (merged) bertingkat.

Header pada workbook BGG membentang beberapa baris dan memakai merged cell,
sehingga `header=0` selalu salah. Resolver di sini:

1. Membuka sheet lewat openpyxl dan MENGEMBANGKAN setiap merged range menjadi
   nilai penuh di seluruh selnya. Ini eksak - tidak menebak lewat forward-fill,
   yang akan ikut mengisi kolom kosong di luar rentang merge yang sebenarnya.
2. Menemukan baris awal data dengan strategi yang dinyatakan per sheet.
3. Meratakan blok header menjadi satu nama per kolom, lalu memetakannya ke nama
   kanonik lewat alias.

Peta hasil resolusi dicetak di audit dan harus dikonfirmasi manusia sebelum
estimasi dijalankan.

Alternatif yang dipertimbangkan dan ditolak:
  - `pandas.read_excel(header=[6,7,8])`: menuntut nomor baris header dihardcode
    per sheet, dan diam-diam salah kalau tata letak bergeser satu baris.
  - Forward-fill horizontal atas hasil pandas: mendekati benar, tapi mengisi
    melewati ujung merge sehingga kolom tak berjudul mewarisi judul tetangga.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

import numpy as np
import openpyxl
import pandas as pd

from ..errors import SchemaError
from ..logging_setup import get_logger

log = get_logger("io.excel")

# Baris pemisah dekoratif pada sheet SLL: seluruhnya "-" atau kosong.
_SEPARATOR = re.compile(r"^[-_\s]*$")
_MIN_HEADER_CELLS = 5
_MIN_DATA_NUMERICS = 3


def _norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def _is_number(value: object) -> bool:
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float, np.integer, np.floating)):
        return not (isinstance(value, float) and np.isnan(value))
    return False


def read_matrix(path: Path, sheet: str) -> list[list[Any]]:
    """Baca sheet menjadi matriks padat, merged range dikembangkan penuh."""
    book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if sheet not in book.sheetnames:
        raise SchemaError(f"sheet '{sheet}' tidak ada di {path.name}. Tersedia: {book.sheetnames}")
    ws = book[sheet]
    matrix = [[cell.value for cell in row] for row in ws.iter_rows()]

    for rng in ws.merged_cells.ranges:
        r0, r1 = rng.min_row - 1, rng.max_row - 1
        c0, c1 = rng.min_col - 1, rng.max_col - 1
        if r0 >= len(matrix) or c0 >= len(matrix[r0]):
            continue
        value = matrix[r0][c0]
        for r in range(r0, min(r1 + 1, len(matrix))):
            for c in range(c0, min(c1 + 1, len(matrix[r]))):
                matrix[r][c] = value
    book.close()
    return matrix


def _row_non_empty(row: Iterable[Any]) -> int:
    return sum(1 for v in row if v is not None and str(v).strip() != "")


def _row_numerics(row: Iterable[Any]) -> int:
    return sum(1 for v in row if _is_number(v))


def _is_separator(row: Iterable[Any]) -> bool:
    values = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]
    return bool(values) and all(_SEPARATOR.match(v) for v in values)


def find_header_block(matrix: list[list[Any]], anchor: str) -> list[int]:
    """Blok header = baris ber-anchor plus baris padat di bawahnya yang bebas angka.

    Baris header berisi teks saja; baris data selalu membawa angka (kedalaman,
    koordinat). Perbedaan itu yang dipakai untuk menutup blok header, bukan
    nomor baris yang dihardcode.

    Alternatif yang ditolak: 'baris data = baris pertama dengan >=3 angka'.
    Baris data SLL hanya membawa dua angka (Depth From dan Depth To), sehingga
    ambang itu melompati delapan baris data pertama dan menariknya menjadi
    header - persis kegagalan yang ditemukan saat resolver diuji ke DH09_05C1.
    """
    target = _norm(anchor)
    anchor_row = None
    for idx, row in enumerate(matrix):
        if any(_norm(v) == target for v in row if v is not None):
            anchor_row = idx
            break
    if anchor_row is None:
        for idx, row in enumerate(matrix):
            if any(target in _norm(v) for v in row if v is not None):
                anchor_row = idx
                break
    if anchor_row is None:
        raise SchemaError(f"baris header tidak ditemukan: anchor '{anchor}' tidak ada di sheet")

    rows = [anchor_row]
    for idx in range(anchor_row + 1, len(matrix)):
        row = matrix[idx]
        if _is_separator(row):
            break
        if _row_numerics(row) > 0 or _row_non_empty(row) < 3:
            break
        rows.append(idx)
    return rows


def find_data_start(matrix: list[list[Any]], header_rows: list[int]) -> int:
    """Baris data pertama setelah blok header: baris pertama yang membawa angka."""
    for idx in range(header_rows[-1] + 1, len(matrix)):
        row = matrix[idx]
        if _is_separator(row):
            continue
        if _row_numerics(row) > 0:
            return idx
    raise SchemaError("tidak ditemukan baris data di bawah blok header")


def find_data_start_after_header(matrix: list[list[Any]]) -> int:
    """Untuk tabel lookup yang isinya teks: data mulai tepat di bawah header."""
    for idx, row in enumerate(matrix):
        if _row_non_empty(row) >= _MIN_HEADER_CELLS:
            return idx + 1
    raise SchemaError("tidak ditemukan baris header pada sheet lookup")


def flatten_header(matrix: list[list[Any]], header_rows: list[int], width: int) -> list[str]:
    """Ratakan blok header menjadi satu nama per kolom.

    Nilai yang berulang secara vertikal (akibat merge) dibuang agar
    'Depth | Depth | From' menjadi 'Depth From'.
    """
    names: list[str] = []
    for col in range(width):
        parts: list[str] = []
        for row in header_rows:
            if col >= len(matrix[row]):
                continue
            value = matrix[row][col]
            if value is None:
                continue
            text = " ".join(str(value).split())
            if text and (not parts or parts[-1] != text):
                parts.append(text)
        names.append(" ".join(parts).strip())
    return names


@dataclass
class HeaderMap:
    """Hasil resolusi header untuk satu sheet."""

    sheet: str
    data_start_row: int          # 0-based
    header_rows: list[int]       # 0-based
    flattened: list[str]
    canonical: dict[str, int]    # nama kanonik -> indeks kolom
    unmapped: list[str] = field(default_factory=list)

    def describe(self) -> str:
        lines = [
            f"  sheet '{self.sheet}': header baris "
            f"{[r + 1 for r in self.header_rows]}, data mulai baris {self.data_start_row + 1}"
        ]
        for name, col in sorted(self.canonical.items(), key=lambda kv: kv[1]):
            lines.append(f"      {name:<22} <- kolom {col:>2}  '{self.flattened[col]}'")
        return "\n".join(lines)


@dataclass
class SheetSpec:
    """Deskripsi deklaratif satu sheet: alias kolom dan cara menemukan data.

    `anchor` adalah teks yang pasti muncul di baris teratas blok header sheet
    ini. Ia dinyatakan per sheet, bukan ditebak, sehingga peta yang dihasilkan
    dapat ditinjau manusia.
    """

    name: str
    aliases: dict[str, list[str]]
    required: list[str]
    anchor: str = "Hole_ID"
    lookup_table: bool = False


def _match_column(flattened: list[str], aliases: list[str]) -> int | None:
    """Cocokkan alias terhadap header yang sudah diratakan.

    Dicoba berurutan: sama persis, lalu berakhiran, lalu mengandung. Urutan ini
    mencegah 'From' pada 'Core Loss From' merebut kolom 'Depth From'.
    """
    normalised = [_norm(name) for name in flattened]
    targets = [_norm(a) for a in aliases]
    for target in targets:
        for idx, name in enumerate(normalised):
            if name == target:
                return idx
    for target in targets:
        for idx, name in enumerate(normalised):
            if name.endswith(target) and name:
                return idx
    for target in targets:
        for idx, name in enumerate(normalised):
            if target and target in name:
                return idx
    return None


@dataclass
class SheetTable:
    header: HeaderMap
    frame: pd.DataFrame
    raw_rows: int


def resolve_sheet(path: Path, spec: SheetSpec) -> SheetTable:
    matrix = read_matrix(path, spec.name)
    width = max((len(r) for r in matrix), default=0)
    matrix = [list(r) + [None] * (width - len(r)) for r in matrix]

    if spec.lookup_table:
        data_start = find_data_start_after_header(matrix)
        header_rows = [data_start - 1]
    else:
        header_rows = find_header_block(matrix, spec.anchor)
        data_start = find_data_start(matrix, header_rows)
    flattened = flatten_header(matrix, header_rows, width)

    canonical: dict[str, int] = {}
    for name, aliases in spec.aliases.items():
        col = _match_column(flattened, aliases)
        if col is not None and col not in canonical.values():
            canonical[name] = col

    missing = [c for c in spec.required if c not in canonical]
    if missing:
        raise SchemaError(
            f"{path.name} / sheet '{spec.name}': kolom wajib tidak terpetakan: {missing}\n"
            f"  header yang teratasi: {[f for f in flattened if f]}"
        )

    body = matrix[data_start:]
    frame = pd.DataFrame(body, columns=range(width))
    frame = frame.rename(columns={col: name for name, col in canonical.items()})
    frame = frame[[c for c in canonical]]
    frame = frame.dropna(how="all").reset_index(drop=True)

    unmapped = [f for i, f in enumerate(flattened) if f and i not in canonical.values()]
    return SheetTable(
        header=HeaderMap(spec.name, data_start, header_rows, flattened, canonical, unmapped),
        frame=frame,
        raw_rows=len(body),
    )


# --------------------------------------------------------------------------- #
# Spesifikasi sheet workbook BGG
# --------------------------------------------------------------------------- #
SLL_ALIASES = {
    "hole_id": ["Hole_ID", "Hole ID"],
    "depth_from": ["Drill Detail Depth From", "Depth From"],
    "depth_to": ["Drill Detail Depth To", "Depth To"],
    "lithology": ["Lithology Lith", "Lithologi", "Lith"],
    "lith_qualifier": ["Lithology Lith Qualifier", "Lith Qualifier"],
    "core_recovery": ["Drill Detail Core Recov", "Core Recov"],
    "core_length": ["Drill Detail Core Length", "Core Length"],
    "hole_type": ["Drill Detail Hole Type", "Drill Detail Continuity", "Hole Type"],
    "log_basis": ["Drill Detail Log Basis", "Log Basis"],
    "seam": ["Lithology Seam", "Seam"],
    "ply": ["Lithology Ply", "Ply"],
}

SPEC_SLL_RECONCILED = SheetSpec("SLL_Reconciled", SLL_ALIASES,
                                ["hole_id", "depth_from", "depth_to", "lithology", "seam"],
                                anchor="Hole_ID")
SPEC_SLL_WELLSITE = SheetSpec("SLL_Wellsite", SLL_ALIASES,
                              ["hole_id", "depth_from", "depth_to", "lithology", "seam"],
                              anchor="Hole_ID")

SPEC_COLLAR = SheetSpec(
    "Collar",
    {
        "hole_id": ["Hole ID"],
        "east": ["Collar Coordinate East", "Coordinate East", "East"],
        "north": ["Collar Coordinate North", "Coordinate North", "North"],
        "rl": ["Collar Coordinate RL", "Coordinate RL", "RL"],
        "survey_method": ["Collar Survey Method", "Survey Method"],
        "survey_company": ["Collar Survey Company", "Survey Company"],
        "max_depth_drilling": ["Drilling Max Depth", "Max Depth"],
        "log_company": ["Geophysical Company", "Company"],
        "log_gm": ["Geophysical Geophysical Log GM", "Geophysical Log GM", "GM"],
        "log_dn": ["Geophysical Geophysical Log DN", "Geophysical Log DN", "DN"],
        "log_cal": ["Geophysical Geophysical Log CAL", "Geophysical Log CAL", "CAL"],
        "log_rs": ["Geophysical Geophysical Log RS", "Geophysical Log RS", "RS"],
        "log_max_depth": ["Geophysical Max Depth", "Etc Max Depth"],
    },
    ["hole_id", "east", "north", "rl"],
    anchor="Hole ID",
)

SPEC_BHC = SheetSpec(
    "BHC",
    {
        "hole_number_actual": ["HOLE NUMBER ACTUAL", "ACTUAL"],
        "hole_number_plan": ["HOLE NUMBER PLAN", "PLAN"],
        "gps_east": ["COORDINATE BY GPS EASTING", "BY GPS EASTING"],
        "gps_north": ["COORDINATE BY GPS NORTHING", "BY GPS NORTHING"],
        "gps_rl": ["COORDINATE BY GPS ELV.", "BY GPS ELV."],
        "ts_east": ["COORDINATE BY TS EASTING", "BY TS EASTING"],
        "ts_north": ["COORDINATE BY TS NORTHING", "BY TS NORTHING"],
        "ts_rl": ["COORDINATE BY TS ELV.", "BY TS ELV."],
        "total_depth_drilling": ["TOTAL DEPTH (m) DRILLING", "DEPTH (m) DRILLING"],
        "total_depth_logging": ["TOTAL DEPTH (m) LOGGING", "DEPTH (m) LOGGING"],
        "core_length": ["CORE LENGTH ( M )"],
        "core_recovery": ["CORE REC. ( % )"],
        "reconcile_from": ["RECONCILE FROM"],
        "reconcile_to": ["RECONCILE TO"],
        "reconcile_thick": ["RECONCILE THICK"],
        "picking_from": ["PICKING FROM"],
        "seam": ["SEAM"],
        "coal_recovery": ["COAL REC ( % )", "COAL REC"],
    },
    ["hole_number_actual"],
    anchor="NO",
)

SPEC_SAMPLING = SheetSpec(
    "Sampling",
    {
        "hole_id": ["Hole Id"],
        "sample_number": ["Sample Number"],
        "sample_position": ["Sample Position"],
        "drilling_from": ["Drilling Depth From"],
        "drilling_to": ["Drilling Depth To"],
        "drilling_length": ["Drilling Depth Length"],
        "adjusted_from": ["Adjusted Depth From"],
        "adjusted_to": ["Adjusted Depth To"],
        "adjusted_length": ["Adjusted Depth Length"],
        "sampling_date": ["Sampling Date"],
        "seam": ["Seam"],
        "remarks": ["Remarks"],
    },
    ["hole_id", "sample_number", "adjusted_from", "adjusted_to"],
    anchor="Sample Number",
)

SPEC_LIBRARY = SheetSpec("Library SLL", {"raw": ["Shift"]}, [], lookup_table=True)


@dataclass
class Workbook:
    """Satu workbook lubang bor yang sudah teratasi headernya."""

    path: Path
    hole_id: str
    sheets: dict[str, SheetTable]
    lithology_library: dict[str, str]

    @property
    def headers(self) -> list[HeaderMap]:
        return [t.header for t in self.sheets.values()]


def parse_lithology_library(path: Path) -> dict[str, str]:
    """Baca tabel kode litologi dari sheet 'Library SLL'.

    Kode TIDAK dihardcode di kode program. Blok 'Lithology' pada sheet ini
    berpasangan kode-deskripsi; keduanya dibaca apa adanya.
    """
    matrix = read_matrix(path, "Library SLL")
    header_row = next(
        (i for i, r in enumerate(matrix) if any(_norm(v) == "lithology" for v in r)), None
    )
    if header_row is None:
        raise SchemaError(f"{path.name}: blok 'Lithology' tidak ditemukan di 'Library SLL'")
    col = next(i for i, v in enumerate(matrix[header_row]) if _norm(v) == "lithology")

    library: dict[str, str] = {}
    for row in matrix[header_row + 1:]:
        if col + 1 >= len(row):
            continue
        code, description = row[col], row[col + 1]
        if code is None or description is None:
            continue
        code_text = str(code).strip()
        if code_text and len(code_text) <= 4:
            library[code_text.upper()] = str(description).strip()
    if not library:
        raise SchemaError(f"{path.name}: tabel kode litologi kosong")
    return library


def load_workbook(path: Path) -> Workbook:
    path = Path(path)
    sheets: dict[str, SheetTable] = {}
    for spec in (SPEC_COLLAR, SPEC_SLL_RECONCILED, SPEC_SLL_WELLSITE,
                 SPEC_BHC, SPEC_SAMPLING):
        try:
            sheets[spec.name] = resolve_sheet(path, spec)
        except SchemaError as exc:
            log.warning(f"{path.name}: sheet '{spec.name}' tidak dapat diresolusi -> {exc}")

    if "Collar" not in sheets:
        raise SchemaError(f"{path.name}: sheet 'Collar' wajib ada dan tidak dapat diresolusi")

    collar = sheets["Collar"].frame
    hole_ids = [str(v).strip() for v in collar["hole_id"].dropna() if str(v).strip()]
    if not hole_ids:
        raise SchemaError(f"{path.name}: sheet 'Collar' tidak memuat hole_id")

    return Workbook(
        path=path,
        hole_id=hole_ids[0],
        sheets=sheets,
        lithology_library=parse_lithology_library(path),
    )


def normalise_hole_id(value: object) -> str:
    """Samakan varian penulisan ID lubang lintas sumber.

    LAS menulis 'DH-09-05C1', workbook 'DH09_05C1', BHC kolom PLAN 'DH09_05C'.
    Tanpa normalisasi, join antar-sumber gagal diam-diam dan lubang hilang.
    """
    return re.sub(r"[^A-Z0-9]", "", str(value).strip().upper())
